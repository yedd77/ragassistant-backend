from io import BytesIO
from pathlib import Path
import fitz
from docx import Document
from pptx import Presentation
from app.config import get_settings

ALLOWED = {".pdf", ".docx", ".pptx", ".txt", ".md"}

class DocumentParseError(ValueError):
    pass

def extract_sections(filename: str, content: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise DocumentParseError("Supported formats: PDF, DOCX, PPTX, TXT, MD.")
    settings = get_settings()
    sections = []
    try:
        if ext == ".pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                for n, page in enumerate(pdf, 1):
                    if n > settings.max_pages:
                        raise DocumentParseError(
                            f"PDF exceeds the {settings.max_pages}-page limit."
                        )
                    text = page.get_text("text").strip()
                    if text:
                        sections.append({"page": n, "text": text})
        elif ext == ".docx":
            doc = Document(BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            if parts:
                sections.append({"page": None, "text": "\n".join(parts)})
        elif ext == ".pptx":
            pres = Presentation(BytesIO(content))
            for n, slide in enumerate(pres.slides, 1):
                parts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        parts.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            parts.append(" | ".join(c.text.strip() for c in row.cells))
                if parts:
                    sections.append({"page": n, "text": "\n".join(parts)})
        else:
            text = content.decode("utf-8-sig", errors="replace").strip()
            if text:
                sections.append({"page": None, "text": text})
    except Exception as exc:
        raise DocumentParseError(f"Could not parse '{filename}': {exc}") from exc
    if not sections:
        raise DocumentParseError("No extractable text found. Scanned PDF OCR is not included yet.")
    return sections
